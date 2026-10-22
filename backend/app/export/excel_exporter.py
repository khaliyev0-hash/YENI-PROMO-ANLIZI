import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
import pandas as pd
from typing import Dict, Any

def generate_multi_sheet_excel(fact_df: pd.DataFrame, timeline_df: pd.DataFrame, analytics_data: Dict[str, Any]) -> bytes:
    """
    Creates a styled 10-sheet Excel workbook:
    SUMMARY, STORE_ANALYSIS, PRODUCT_ANALYSIS, STOCKOUT_ANALYSIS,
    STOCK_ANALYSIS, FORECAST_ACCURACY, DAILY_STOCK, ALERTS,
    RECOMMENDATIONS, RAW_DATA
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Dark slate
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="0F172A")
    regular_font = Font(name="Calibri", size=11)
    bold_font = Font(name="Calibri", size=11, bold=True)
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    def style_sheet_headers(ws, row_idx=1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border_thin

    def auto_fit_columns(ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # 1. SUMMARY SHEET
    ws_summary = wb.create_sheet(title="SUMMARY")
    ws_summary.cell(row=1, column=1, value="TAMSTORE PROMO INTELLIGENCE - İCRAÇI XÜLASƏ").font = title_font
    
    kpi = analytics_data.get("kpi_cards", {})
    summary_data = [
        ["Metrika", "Dəyər"],
        ["Ümumi Mağaza Sayı", kpi.get("total_stores", 0)],
        ["Ümumi Məhsul Sayı", kpi.get("total_products", 0)],
        ["Ümumi Proqnoz (FC)", kpi.get("total_forecast", 0)],
        ["Ümumi Satış (SALE)", kpi.get("total_sales", 0)],
        ["Proqnoz Reallaşması (%)", f"{kpi.get('achievement_pct', 0)}%"],
        ["Ümumi Qalıq Stok", kpi.get("total_stock", 0)],
        ["Stokout Mağaza Sayı", kpi.get("stockout_stores_count", 0)],
        ["Stokout Məhsul Sayı", kpi.get("stockout_products_count", 0)],
        ["Stokout Faizi (%)", f"{kpi.get('stockout_rate_pct', 0)}%"],
        ["Satışdan Keçmə Faizi (Sell Through %)", f"{kpi.get('sell_through_pct', 0)}%"],
        ["Proqnoz Xətası (WAPE %)", f"{kpi.get('wape_pct', 0)}%"],
        ["Promo Sağlamlıq Balı (0-100)", kpi.get("promo_health_score", 0)],
        ["Promo Sağlamlıq Statusu", kpi.get("promo_health_status", "")]
    ]
    for r_idx, row in enumerate(summary_data, start=3):
        for c_idx, val in enumerate(row, start=1):
            cell = ws_summary.cell(row=r_idx, column=c_idx, value=val)
            cell.font = bold_font if r_idx == 3 else regular_font
            cell.border = border_thin
    style_sheet_headers(ws_summary, row_idx=3)
    auto_fit_columns(ws_summary)

    # 2. STORE_ANALYSIS SHEET
    ws_store = wb.create_sheet(title="STORE_ANALYSIS")
    stores_data = analytics_data.get("store_performances", [])
    if stores_data:
        store_df = pd.DataFrame(stores_data)
        for r in dataframe_to_rows(store_df, index=False, header=True):
            ws_store.append(r)
        style_sheet_headers(ws_store)
        auto_fit_columns(ws_store)

    # 3. PRODUCT_ANALYSIS SHEET
    ws_prod = wb.create_sheet(title="PRODUCT_ANALYSIS")
    products_data = analytics_data.get("product_performances", [])
    if products_data:
        prod_df = pd.DataFrame(products_data)
        for r in dataframe_to_rows(prod_df, index=False, header=True):
            ws_prod.append(r)
        style_sheet_headers(ws_prod)
        auto_fit_columns(ws_prod)

    # 4. STOCKOUT_ANALYSIS SHEET
    ws_stockout = wb.create_sheet(title="STOCKOUT_ANALYSIS")
    ws_stockout.append(["Store", "Code", "Name", "Forecast", "Sales", "Final Stock", "Status"])
    stockout_rows = fact_df[fact_df['FINAL_STOCK'] <= 0.0001]
    for _, row in stockout_rows.iterrows():
        ws_stockout.append([row['STORE'], row['CODE'], row['NAME'], row['FC'], row['SALE'], row['FINAL_STOCK'], "STOCKOUT"])
    style_sheet_headers(ws_stockout)
    auto_fit_columns(ws_stockout)

    # 5. STOCK_ANALYSIS SHEET
    ws_stock = wb.create_sheet(title="STOCK_ANALYSIS")
    ws_stock.append(["Store", "Code", "Name", "Final Stock", "Sell Through %"])
    for _, row in fact_df.iterrows():
        s = float(row['SALE'])
        stk = float(row['FINAL_STOCK'])
        denom = s + stk
        st_pct = round((s / denom) * 100.0, 2) if denom > 0 else 0.0
        ws_stock.append([row['STORE'], row['CODE'], row['NAME'], stk, f"{st_pct}%"])
    style_sheet_headers(ws_stock)
    auto_fit_columns(ws_stock)

    # 6. FORECAST_ACCURACY SHEET
    ws_acc = wb.create_sheet(title="FORECAST_ACCURACY")
    acc = analytics_data.get("accuracy_metrics", {})
    ws_acc.append(["Metrika", "Dəyər", "İzah"])
    ws_acc.append(["WAPE (%)", acc.get("wape", 0), "Weighted Absolute Percentage Error - Aşağı olması dəqiqliyi göstərir"])
    ws_acc.append(["MAE", acc.get("mae", 0), "Mean Absolute Error"])
    ws_acc.append(["RMSE", acc.get("rmse", 0), "Root Mean Squared Error"])
    ws_acc.append(["Bias", acc.get("bias", 0), "Ümumi Meyllilik - Müsbət = Tələb çox olub, Mənfi = Aşırı proqnoz"])
    style_sheet_headers(ws_acc)
    auto_fit_columns(ws_acc)

    # 7. DAILY_STOCK SHEET
    ws_daily = wb.create_sheet(title="DAILY_STOCK")
    if not timeline_df.empty:
        for r in dataframe_to_rows(timeline_df, index=False, header=True):
            ws_daily.append(r)
        style_sheet_headers(ws_daily)
        auto_fit_columns(ws_daily)

    # 8. ALERTS SHEET
    ws_alerts = wb.create_sheet(title="ALERTS")
    alerts_data = analytics_data.get("alerts", [])
    if alerts_data:
        al_df = pd.DataFrame(alerts_data)
        for r in dataframe_to_rows(al_df, index=False, header=True):
            ws_alerts.append(r)
        style_sheet_headers(ws_alerts)
        auto_fit_columns(ws_alerts)

    # 9. RECOMMENDATIONS SHEET
    ws_recs = wb.create_sheet(title="RECOMMENDATIONS")
    recs_data = analytics_data.get("recommendations", [])
    if recs_data:
        rec_df = pd.DataFrame(recs_data)
        for r in dataframe_to_rows(rec_df, index=False, header=True):
            ws_recs.append(r)
        style_sheet_headers(ws_recs)
        auto_fit_columns(ws_recs)

    # 10. RAW_DATA SHEET
    ws_raw = wb.create_sheet(title="RAW_DATA")
    for r in dataframe_to_rows(fact_df, index=False, header=True):
        ws_raw.append(r)
    style_sheet_headers(ws_raw)
    auto_fit_columns(ws_raw)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
