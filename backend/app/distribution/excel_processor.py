import io
from datetime import datetime
from typing import Dict, Any, List, Tuple, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel

class ImportedRow(BaseModel):
    row_number: int
    store_code: str
    product_code: str
    product_name: str
    shelf_capacity: float

class RejectedRow(BaseModel):
    row_number: int
    raw_store_code: Optional[str] = None
    raw_product_code: Optional[str] = None
    raw_product_name: Optional[str] = None
    raw_shelf_capacity: Optional[str] = None
    rejection_reason: str

class ImportValidationResult(BaseModel):
    total_rows: int
    valid_count: int
    rejected_count: int
    duplicate_count: int
    distinct_stores: int
    distinct_products: int
    valid_rows: List[ImportedRow]
    rejected_rows: List[RejectedRow]

def normalize_code_string(val: Any) -> str:
    """Preserves leading zeros and formats codes cleanly as strings."""
    if val is None:
        return ""
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val)).strip()
        return str(val).strip()
    s = str(val).strip()
    # If Excel numeric string like '101.0'
    if s.endswith(".0"):
        s = s[:-2]
    return s

def parse_and_validate_excel_template(file_bytes: bytes) -> ImportValidationResult:
    """
    Parses and validates the 4-column promo distribution Excel template.
    Columns required:
    - StoreCode (or MAGAZA, IS YERI NO)
    - ProductCode (or Mehsul_kodu, URUN KODU)
    - ProductName (or Mehsul_adi, URUN AD)
    - ShelfCapacity (or Ref, REF, Tutumu)
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        return ImportValidationResult(
            total_rows=0, valid_count=0, rejected_count=0, duplicate_count=0,
            distinct_stores=0, distinct_products=0, valid_rows=[], rejected_rows=[]
        )

    # Detect header columns
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    header_lower = [c.lower() for c in header]

    store_idx = -1
    product_idx = -1
    name_idx = -1
    shelf_idx = -1

    for idx, col in enumerate(header_lower):
        if col in ["storecode", "store_code", "magaza", "magaza kodu", "is yeri no", "is_yeri_no", "store"]:
            if store_idx == -1:
                store_idx = idx
        elif col in ["productcode", "product_code", "mehsul_kodu", "mehsul kodu", "urun kodu", "urun_kodu", "sku", "item_code"]:
            if product_idx == -1:
                product_idx = idx
        elif col in ["productname", "product_name", "mehsul_adi", "mehsul adi", "urun ad", "urun_adi", "name", "item_name"]:
            if name_idx == -1:
                name_idx = idx
        elif col in ["shelfcapacity", "shelf_capacity", "ref", "rəf", "ref tutumu", "tutumu", "capacity"]:
            if shelf_idx == -1:
                shelf_idx = idx

    # If canonical columns not explicitly named, fallback to standard column order (0: Store, 1: SKU, 2: Name, 3: Shelf)
    if store_idx == -1 and len(header) >= 1:
        store_idx = 0
    if product_idx == -1 and len(header) >= 2:
        product_idx = 1
    if name_idx == -1 and len(header) >= 3:
        name_idx = 2
    if shelf_idx == -1 and len(header) >= 4:
        shelf_idx = 3

    valid_rows: List[ImportedRow] = []
    rejected_rows: List[RejectedRow] = []
    seen_combinations = set()
    duplicate_count = 0

    for row_idx, r in enumerate(rows[1:], start=2):
        if not r or all(c is None or str(c).strip() == "" for c in r):
            continue  # Skip completely empty rows

        raw_store = r[store_idx] if store_idx < len(r) else None
        raw_prod = r[product_idx] if product_idx < len(r) else None
        raw_name = r[name_idx] if name_idx < len(r) else None
        raw_shelf = r[shelf_idx] if shelf_idx < len(r) else None

        st_code = normalize_code_string(raw_store)
        p_code = normalize_code_string(raw_prod)
        p_name = str(raw_name).strip() if raw_name is not None else ""
        
        # Validation checks
        if not st_code:
            rejected_rows.append(RejectedRow(
                row_number=row_idx,
                raw_store_code=str(raw_store),
                raw_product_code=str(raw_prod),
                raw_product_name=p_name,
                raw_shelf_capacity=str(raw_shelf),
                rejection_reason="Mağaza kodu boşdur (StoreCode is empty)"
            ))
            continue

        if not p_code:
            rejected_rows.append(RejectedRow(
                row_number=row_idx,
                raw_store_code=st_code,
                raw_product_code=str(raw_prod),
                raw_product_name=p_name,
                raw_shelf_capacity=str(raw_shelf),
                rejection_reason="Məhsul kodu boşdur (ProductCode is empty)"
            ))
            continue

        # Validate shelf capacity
        shelf_cap = 0.0
        if raw_shelf is not None and str(raw_shelf).strip() != "":
            try:
                shelf_cap = float(raw_shelf)
                if shelf_cap < 0:
                    rejected_rows.append(RejectedRow(
                        row_number=row_idx,
                        raw_store_code=st_code,
                        raw_product_code=p_code,
                        raw_product_name=p_name,
                        raw_shelf_capacity=str(raw_shelf),
                        rejection_reason="Rəf tutumu mənfi ola bilməz (ShelfCapacity < 0)"
                    ))
                    continue
            except (ValueError, TypeError):
                rejected_rows.append(RejectedRow(
                    row_number=row_idx,
                    raw_store_code=st_code,
                    raw_product_code=p_code,
                    raw_product_name=p_name,
                    raw_shelf_capacity=str(raw_shelf),
                    rejection_reason=f"Etibarsız rəf tutumu formatı ({raw_shelf})"
                ))
                continue

        # Duplicate check within Excel file
        pair_key = (st_code, p_code)
        if pair_key in seen_combinations:
            duplicate_count += 1
            rejected_rows.append(RejectedRow(
                row_number=row_idx,
                raw_store_code=st_code,
                raw_product_code=p_code,
                raw_product_name=p_name,
                raw_shelf_capacity=str(raw_shelf),
                rejection_reason="Təkrarlanan mağaza-məhsul cütü (Duplicate Store-Product combination in Excel)"
            ))
            continue

        seen_combinations.add(pair_key)
        valid_rows.append(ImportedRow(
            row_number=row_idx,
            store_code=st_code,
            product_code=p_code,
            product_name=p_name if p_name else p_code,
            shelf_capacity=shelf_cap
        ))

    distinct_stores = len(set(r.store_code for r in valid_rows))
    distinct_products = len(set(r.product_code for r in valid_rows))

    return ImportValidationResult(
        total_rows=len(valid_rows) + len(rejected_rows),
        valid_count=len(valid_rows),
        rejected_count=len(rejected_rows),
        duplicate_count=duplicate_count,
        distinct_stores=distinct_stores,
        distinct_products=distinct_products,
        valid_rows=valid_rows,
        rejected_rows=rejected_rows
    )

def generate_rejected_rows_excel(rejected_rows: List[RejectedRow]) -> bytes:
    """Generates an Excel file containing all rejected rows and reasons for user download."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rejected_Rows"

    headers = ["Sətr No", "Mağaza Kodu", "Məhsul Kodu", "Məhsul Adı", "Rəf Tutumu", "İmtina Səbəbi"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in rejected_rows:
        ws.append([
            r.row_number,
            r.raw_store_code or "",
            r.raw_product_code or "",
            r.raw_product_name or "",
            r.raw_shelf_capacity or "",
            r.rejection_reason
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def generate_distribution_6sheet_excel(
    summary_data: Dict[str, Any],
    orders_data: List[Dict[str, Any]],
    store_summary: List[Dict[str, Any]],
    sku_summary: List[Dict[str, Any]],
    exceptions_data: List[Dict[str, Any]],
    parameters_data: Dict[str, Any]
) -> bytes:
    """
    Generates the complete 6-sheet professional workbook required by Master Prompt Section 12:
    - Sheet 1: Summary
    - Sheet 2: Store_SKU_Orders
    - Sheet 3: Store_Summary
    - Sheet 4: SKU_Summary
    - Sheet 5: Exceptions
    - Sheet 6: Parameters
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    navy_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    blue_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    warn_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # ----------------------------------------------------
    # Sheet 1: Summary
    # ----------------------------------------------------
    ws1 = wb.create_sheet(title="Summary")
    ws1.append(["PROMO DISTRIBUTION OPTIMIZATION — SUMMARY"])
    ws1.cell(row=1, column=1).font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    ws1.append([])

    conf_summary = summary_data.get("confidence_summary", {})
    conf_str = f"NORMAL: {conf_summary.get('NORMAL', 0)} | LIMITED: {conf_summary.get('LIMITED', 0)} | MANUAL_REVIEW: {conf_summary.get('MANUAL_REVIEW', 0)}"

    kpis = [
        ("Kampaniya Başlanğıc Tarixi", summary_data.get("start_date")),
        ("Kampaniya Bitiş Tarixi", summary_data.get("end_date")),
        ("Kampaniya Müddəti (Gün)", summary_data.get("duration_days")),
        ("Promo Uplift Artımı (%)", f"+{summary_data.get('promo_uplift_pct', 0)}%"),
        ("Stok-out Düzəlişi", "Deaktiv (Tarixi stok məlumatı olmadığı üçün təhlükəsiz proqnoz: 30 təqvim gününə bölünmə)"),
        ("Tarixi Stok Mövcudluğu", "Yoxdur (Yalnız cari anbar qalığı mövcuddur)"),
        ("Keçən İlin Eyni Dövrü (YoY)", f"Aktiv ({int(summary_data.get('yoy_weight', 0.3)*100)}% çəki, {summary_data.get('yoy_years_back', 1)} il əvvəl)" if summary_data.get("include_yoy_history", True) else "Deaktiv"),
        ("Proqnoz Dərəcələri (Confidence)", conf_str),
        ("Mağaza Sayı", summary_data.get("store_count")),
        ("Məhsul Sayı", summary_data.get("product_count")),
        ("Ümumi Proqnoz Tələbat (Ədəd)", summary_data.get("total_forecast_demand")),
        ("Ümumi Cari Stok (Ədəd)", summary_data.get("total_current_stock")),
        ("Ümumi Xalis Ehtiyac (Ədəd)", summary_data.get("total_net_requirement")),
        ("Ümumi Tövsiyə Olunan Sifariş (Ədəd)", summary_data.get("total_recommended_order_units")),
        ("Ümumi Tövsiyə Olunan Qutu Sayı", summary_data.get("total_recommended_cases")),
        ("Yuvarlaqlaşdırmadan Artıq Miqdar (Ədəd)", summary_data.get("total_excess_units")),
        ("Xəbərdarlıq və İstisna Sayı", summary_data.get("warning_count")),
        ("Hesablama Vaxtı", summary_data.get("created_at"))
    ]

    for label, val in kpis:
        ws1.append([label, val])

    ws1.column_dimensions['A'].width = 38
    ws1.column_dimensions['B'].width = 35
    for r_idx in range(3, 3 + len(kpis)):
        ws1.cell(row=r_idx, column=1).font = Font(name="Calibri", size=11, bold=True)
        ws1.cell(row=r_idx, column=1).fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        ws1.cell(row=r_idx, column=2).alignment = Alignment(horizontal="right")

    # ----------------------------------------------------
    # Sheet 2: Store_SKU_Orders
    # ----------------------------------------------------
    ws2 = wb.create_sheet(title="Store_SKU_Orders")
    s2_cols = [
        "StoreCode", "StoreName", "ProductCode", "ProductName",
        "PromotionStartDate", "PromotionEndDate", "DurationDays",
        "RecentTrendDemand", "YoYHistoricalSales", "YoYPromoSales",
        "ForecastDemand", "ForecastMethod", "ForecastConfidence",
        "StockoutCorrectionApplied", "StockoutHistoryAvailable", "YoYDataAvailable",
        "PromoUpliftSource",
        "CurrentStock", "IncomingStock", "TargetEndingStock",
        "ShelfCapacity", "NetRequirement", "RecommendedOrderUnits",
        "CaseQuantity", "RecommendedCases", "BlockQuantity", "RecommendedBlocks",
        "ExpectedEndingStock", "ExcessUnitsFromRounding", "WarningStatus", "WarningReason"
    ]
    ws2.append(s2_cols)
    for c_idx in range(1, len(s2_cols) + 1):
        cell = ws2.cell(row=1, column=c_idx)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for o in orders_data:
        ws2.append([
            str(o.get("store_code") or ""),
            str(o.get("store_name") or ""),
            str(o.get("product_code") or ""),
            str(o.get("product_name") or ""),
            str(o.get("start_date") or ""),
            str(o.get("end_date") or ""),
            o.get("duration_days"),
            o.get("recent_trend_demand"),
            o.get("yoy_historical_sales"),
            o.get("yoy_promo_sales"),
            o.get("forecast_demand"),
            str(o.get("forecast_method") or "CURRENT_TREND_ONLY"),
            str(o.get("forecast_confidence") or "NORMAL"),
            bool(o.get("stockout_correction_applied", False)),
            bool(o.get("stockout_history_available", False)),
            bool(o.get("yoy_data_available", False)),
            str(o.get("promo_uplift_source") or "NONE"),
            o.get("current_stock"),
            o.get("eligible_incoming_stock"),
            o.get("target_ending_stock"),
            o.get("shelf_capacity"),
            o.get("net_requirement"),
            o.get("recommended_order_units"),
            o.get("case_quantity"),
            o.get("recommended_cases"),
            o.get("block_quantity"),
            o.get("recommended_blocks"),
            o.get("expected_ending_stock"),
            o.get("excess_units_from_rounding"),
            str(o.get("warning_status") or "OK"),
            str(o.get("warning_reason") or o.get("warning_message") or "")
        ])

    for col in ws2.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 11)

    # ----------------------------------------------------
    # Sheet 3: Store_Summary
    # ----------------------------------------------------
    ws3 = wb.create_sheet(title="Store_Summary")
    s3_cols = ["StoreCode", "StoreName", "ProductCount", "TotalForecastDemand", "TotalCurrentStock", "TotalRecommendedOrder", "WarningCount"]
    ws3.append(s3_cols)
    for c_idx in range(1, len(s3_cols) + 1):
        cell = ws3.cell(row=1, column=c_idx)
        cell.fill = blue_fill
        cell.font = header_font

    for st in store_summary:
        ws3.append([
            str(st.get("store_code") or ""),
            str(st.get("store_name") or ""),
            st.get("product_count"),
            st.get("total_forecast_demand"),
            st.get("total_current_stock"),
            st.get("total_recommended_order"),
            st.get("warning_count")
        ])
    for col in ws3.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws3.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 12)

    # ----------------------------------------------------
    # Sheet 4: SKU_Summary
    # ----------------------------------------------------
    ws4 = wb.create_sheet(title="SKU_Summary")
    s4_cols = ["ProductCode", "ProductName", "StoreCount", "TotalForecastDemand", "TotalCurrentStock", "TotalRecommendedOrder", "WarningCount"]
    ws4.append(s4_cols)
    for c_idx in range(1, len(s4_cols) + 1):
        cell = ws4.cell(row=1, column=c_idx)
        cell.fill = blue_fill
        cell.font = header_font

    for sku in sku_summary:
        ws4.append([
            str(sku.get("product_code") or ""),
            str(sku.get("product_name") or ""),
            sku.get("store_count"),
            sku.get("total_forecast_demand"),
            sku.get("total_current_stock"),
            sku.get("total_recommended_order"),
            sku.get("warning_count")
        ])
    for col in ws4.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws4.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 12)

    # ----------------------------------------------------
    # Sheet 5: Exceptions
    # ----------------------------------------------------
    ws5 = wb.create_sheet(title="Exceptions")
    s5_cols = ["StoreCode", "ProductCode", "ProblemType", "ProblemDescription", "Source", "AffectedCalculations", "ProposedResolution"]
    ws5.append(s5_cols)
    for c_idx in range(1, len(s5_cols) + 1):
        cell = ws5.cell(row=1, column=c_idx)
        cell.fill = navy_fill
        cell.font = header_font

    for exc in exceptions_data:
        ws5.append([
            str(exc.get("store_code") or ""),
            str(exc.get("product_code") or ""),
            str(exc.get("exception_type") or ""),
            str(exc.get("description") or ""),
            str(exc.get("source") or ""),
            str(exc.get("affected_calculations") or ""),
            str(exc.get("proposed_resolution") or "")
        ])
    for col in ws5.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws5.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 14)

    # ----------------------------------------------------
    # Sheet 6: Parameters
    # ----------------------------------------------------
    ws6 = wb.create_sheet(title="Parameters")
    ws6.append(["Planlaşdırma Parametri", "Dəyər"])
    ws6.cell(row=1, column=1).font = header_font
    ws6.cell(row=1, column=1).fill = navy_fill
    ws6.cell(row=1, column=2).font = header_font
    ws6.cell(row=1, column=2).fill = navy_fill

    for k, v in parameters_data.items():
        ws6.append([k, str(v)])

    ws6.column_dimensions['A'].width = 35
    ws6.column_dimensions['B'].width = 50

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
