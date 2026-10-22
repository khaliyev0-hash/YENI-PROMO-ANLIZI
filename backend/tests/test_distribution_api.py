import io
import pytest
from unittest.mock import patch
from fastapi.testclient import TestClient
from app.main import app
import openpyxl

client = TestClient(app)

def create_sample_excel_bytes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["StoreCode", "ProductCode", "ProductName", "ShelfCapacity"])
    ws.append(["101", "100003", "AHMAD TEA 100G", 24])
    ws.append(["102", "100003", "AHMAD TEA 100G", 24])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def test_api_health():
    res = client.get("/api/health")
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "healthy"

def test_api_source_status():
    res = client.get("/api/source/status")
    assert res.status_code == 200
    data = res.json()
    assert "status" in data

def test_api_import_preview():
    file_bytes = create_sample_excel_bytes()
    res = client.post(
        "/api/import/preview",
        files={"file": ("test.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert res.status_code == 200
    data = res.json()
    assert data["valid_count"] == 2
    assert data["rejected_count"] == 0
    assert len(data["preview_rows"]) == 2
    assert data["preview_rows"][0]["store_code"] == "101"

def test_api_planning_run_lifecycle():
    # 1. Create a planning run
    mock_raw_stocks = [
        {
            "IS YERI NO": "101", "URUN KODU": "100003", "URUN AD": "AHMAD TEA 100G",
            "IS YERI TEDARUK TIPI": 41, "TEDARİKCİ BLOKAJİ": "XEYR",
            "Stok miqdari": 5.0, "QUTU ICI": 12.0, "BLOK ICI": 1.0, "VAHIDI": "ADET"
        },
        {
            "IS YERI NO": "102", "URUN KODU": "100003", "URUN AD": "AHMAD TEA 100G",
            "IS YERI TEDARUK TIPI": 41, "TEDARİKCİ BLOKAJİ": "XEYR",
            "Stok miqdari": 30.0, "QUTU ICI": 12.0, "BLOK ICI": 1.0, "VAHIDI": "ADET"
        }
    ]
    mock_forecast = {
        ("101", "100003"): 20.0,
        ("102", "100003"): 10.0
    }

    with patch("app.distribution.service.query_stock_details_for_skus", return_value=mock_raw_stocks), \
         patch("app.distribution.service.query_sales_history_and_forecast", return_value=(mock_forecast, {"100003": 15.0})):

        payload = {
            "title": "API Test Run",
            "start_date": "2026-11-01",
            "duration_days": 7,
            "target_ending_stock": 0.0,
            "include_incoming_stock": True,
            "rows": [
                {"row_number": 2, "store_code": "101", "product_code": "100003", "product_name": "AHMAD TEA 100G", "shelf_capacity": 24.0},
                {"row_number": 3, "store_code": "102", "product_code": "100003", "product_name": "AHMAD TEA 100G", "shelf_capacity": 24.0}
            ]
        }
        res_create = client.post("/api/planning-runs", json=payload)
        assert res_create.status_code == 200
        run_data = res_create.json()
        assert run_data["status"] == "success"
        run_id = run_data["run_id"]

        # 2. List planning runs
        res_list = client.get("/api/planning-runs")
        assert res_list.status_code == 200
        runs_list = res_list.json()["runs"]
        assert any(r["id"] == run_id for r in runs_list)

        # 3. Get single planning run
        res_single = client.get(f"/api/planning-runs/{run_id}")
        assert res_single.status_code == 200
        single_data = res_single.json()
        assert single_data["summary"]["store_count"] == 2

        # 4. Get planning results with pagination & filter
        res_results = client.get(f"/api/planning-runs/{run_id}/results?page=1&page_size=10")
        assert res_results.status_code == 200
        res_items = res_results.json()
        assert res_items["total_items"] == 2
        assert len(res_items["items"]) == 2

        # 5. Filter by store
        res_filter = client.get(f"/api/planning-runs/{run_id}/results?store=101")
        assert res_filter.status_code == 200
        assert res_filter.json()["total_items"] == 1

        # 6. Get exceptions
        res_exc = client.get(f"/api/planning-runs/{run_id}/exceptions")
        assert res_exc.status_code == 200

        # 7. Export 6-sheet Excel
        res_export = client.get(f"/api/planning-runs/{run_id}/export")
        assert res_export.status_code == 200
        assert "application/vnd.openxmlformats" in res_export.headers["content-type"]
        wb = openpyxl.load_workbook(io.BytesIO(res_export.content))
        assert "Summary" in wb.sheetnames
        assert "Store_SKU_Orders" in wb.sheetnames
        assert "Store_Summary" in wb.sheetnames
        assert "SKU_Summary" in wb.sheetnames
        assert "Exceptions" in wb.sheetnames
        assert "Parameters" in wb.sheetnames
