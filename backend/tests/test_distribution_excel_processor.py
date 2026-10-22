import io
import pytest
import openpyxl
from app.distribution.excel_processor import (
    parse_and_validate_excel_template,
    generate_rejected_rows_excel,
    generate_distribution_6sheet_excel
)

def create_test_excel_bytes(rows_data):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows_data:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def test_parse_valid_template_preserving_leading_zeros():
    rows = [
        ["StoreCode", "ProductCode", "ProductName", "ShelfCapacity"],
        ["0101", "005432", "SPECIAL TEA", 24],
        ["0102", "100003", "COFFEE PACK", 12]
    ]
    raw_bytes = create_test_excel_bytes(rows)
    res = parse_and_validate_excel_template(raw_bytes)

    assert res.total_rows == 2
    assert res.valid_count == 2
    assert res.rejected_count == 0
    # Leading zeros preserved
    assert res.valid_rows[0].store_code == "0101"
    assert res.valid_rows[0].product_code == "005432"
    assert res.valid_rows[0].shelf_capacity == 24.0

def test_detect_duplicate_and_invalid_rows():
    rows = [
        ["StoreCode", "ProductCode", "ProductName", "ShelfCapacity"],
        ["101", "100003", "AHMAD TEA", 24],
        ["101", "100003", "AHMAD TEA DUPLICATE", 24], # Duplicate
        ["", "100004", "EMPTY STORE", 10],            # Empty store
        ["102", "", "EMPTY PRODUCT", 10],             # Empty product
        ["103", "100005", "NEGATIVE CAPACITY", -5],    # Negative shelf capacity
        ["104", "100006", "TEXT CAPACITY", "INVALID"]  # Invalid shelf capacity
    ]
    raw_bytes = create_test_excel_bytes(rows)
    res = parse_and_validate_excel_template(raw_bytes)

    assert res.valid_count == 1
    assert res.rejected_count == 5
    assert res.duplicate_count == 1

    reasons = [r.rejection_reason for r in res.rejected_rows]
    assert any("Təkrarlanan" in r for r in reasons)
    assert any("Mağaza kodu boşdur" in r for r in reasons)
    assert any("Məhsul kodu boşdur" in r for r in reasons)
    assert any("mənfi ola bilməz" in r for r in reasons)
    assert any("Etibarsız rəf tutumu" in r for r in reasons)

def test_generate_rejected_rows_excel():
    rows = [
        ["StoreCode", "ProductCode", "ProductName", "ShelfCapacity"],
        ["101", "", "ERROR ITEM", 10]
    ]
    raw_bytes = create_test_excel_bytes(rows)
    res = parse_and_validate_excel_template(raw_bytes)
    assert len(res.rejected_rows) == 1

    rejected_excel_bytes = generate_rejected_rows_excel(res.rejected_rows)
    wb = openpyxl.load_workbook(io.BytesIO(rejected_excel_bytes))
    assert "Rejected_Rows" in wb.sheetnames
    ws = wb["Rejected_Rows"]
    assert ws.cell(row=2, column=6).value == "Məhsul kodu boşdur (ProductCode is empty)"

def test_generate_6sheet_excel():
    summary_data = {
        "start_date": "2026-11-01",
        "end_date": "2026-11-07",
        "duration_days": 7,
        "store_count": 1,
        "product_count": 1,
        "total_forecast_demand": 25.0,
        "total_current_stock": 10.0,
        "total_net_requirement": 15.0,
        "total_recommended_order_units": 24.0,
        "total_recommended_cases": 2,
        "total_excess_units": 9.0,
        "warning_count": 0,
        "created_at": "2026-10-21 12:00:00"
    }

    orders_data = [{
        "store_code": "101",
        "store_name": "BAKU STORE",
        "product_code": "100003",
        "product_name": "AHMAD TEA",
        "start_date": "2026-11-01",
        "end_date": "2026-11-07",
        "duration_days": 7,
        "forecast_demand": 25.0,
        "current_stock": 10.0,
        "eligible_incoming_stock": 0.0,
        "target_ending_stock": 0.0,
        "shelf_capacity": 50.0,
        "net_requirement": 15.0,
        "recommended_order_units": 24.0,
        "case_quantity": 12.0,
        "recommended_cases": 2,
        "block_quantity": 1.0,
        "recommended_blocks": 24,
        "expected_ending_stock": 9.0,
        "excess_units_from_rounding": 9.0,
        "warning_status": "OK"
    }]

    store_summary = [{
        "store_code": "101",
        "store_name": "BAKU STORE",
        "product_count": 1,
        "total_forecast_demand": 25.0,
        "total_current_stock": 10.0,
        "total_recommended_order": 24.0,
        "warning_count": 0
    }]

    sku_summary = [{
        "product_code": "100003",
        "product_name": "AHMAD TEA",
        "store_count": 1,
        "total_forecast_demand": 25.0,
        "total_current_stock": 10.0,
        "total_recommended_order": 24.0,
        "warning_count": 0
    }]

    exceptions_data = [{
        "store_code": "101",
        "product_code": "100003",
        "exception_type": "TEST_EXC",
        "description": "Test exception",
        "source": "vw_TS_ITEMS_DETAIL_MAGAZA_FOR_PLANNING",
        "affected_calculations": "None",
        "proposed_resolution": "No action"
    }]

    parameters_data = {
        "Kampaniya Başlanğıc Tarixi": "2026-11-01",
        "Müddət": "7 gün",
        "Proqnoz Metodu": "Baseline ADS & Day-of-Week Pattern",
        "Təchizatçı Blokajı": "Tip 41 + BELI istisna edilir"
    }

    excel_bytes = generate_distribution_6sheet_excel(
        summary_data=summary_data,
        orders_data=orders_data,
        store_summary=store_summary,
        sku_summary=sku_summary,
        exceptions_data=exceptions_data,
        parameters_data=parameters_data
    )

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    # Verify all 6 sheets exist
    assert "Summary" in wb.sheetnames
    assert "Store_SKU_Orders" in wb.sheetnames
    assert "Store_Summary" in wb.sheetnames
    assert "SKU_Summary" in wb.sheetnames
    assert "Exceptions" in wb.sheetnames
    assert "Parameters" in wb.sheetnames

    # Check Sheet 2 content
    ws2 = wb["Store_SKU_Orders"]
    assert ws2.cell(row=2, column=1).value == "101"
    assert ws2.cell(row=2, column=3).value == "100003"
    assert ws2.cell(row=2, column=14).value == 24.0
